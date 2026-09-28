"""Ek usta şablonları (recipes.py'deki seçici kullanır). Her biri gerçekten
çalıştırılarak doğrulandı — bkz. tests/test_recipes.py."""
from __future__ import annotations

TKINTER_HEADLESS = r'''
# GUI KALIBI — Tkinter + otomatik doğrulama için --headless-test modu
import argparse, json
from pathlib import Path

def compute(items: list[str]) -> dict:              # asıl iş GUI'den BAĞIMSIZ bir fonksiyonda
    return {"adet": len(items), "ilk": items[0] if items else None}

def run_gui() -> None:
    import tkinter as tk                              # tkinter yalnızca GUI modunda içe aktarılır
    root = tk.Tk()
    root.title("Uygulama")
    out = tk.Label(root, text="")
    out.pack(padx=12, pady=12)
    tk.Button(root, text="Hesapla", command=lambda: out.config(text=json.dumps(compute(["a", "b"])))).pack()
    root.mainloop()

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--headless-test", action="store_true")
    if ap.parse_args().headless_test:                 # pencere AÇMADAN asıl işi yap, dosyaya yaz, çık
        Path("sonuc.json").write_text(json.dumps(compute(["a", "b"])), encoding="utf-8")
        print("SUCCESS")
    else:
        run_gui()
'''

PILLOW_IMAGES = r'''
# RESİM İŞLEME KALIBI — Pillow, oranı koruyarak küçült, kaynağa dokunma
import sys
from pathlib import Path
from PIL import Image

EXTS = {".jpg", ".jpeg", ".png", ".webp", ".bmp"}

def thumbnails(src: Path, out: Path, size: tuple[int, int] = (256, 256)) -> int:
    out.mkdir(parents=True, exist_ok=True)
    n = 0
    for p in sorted(src.rglob("*")):
        if p.is_file() and p.suffix.lower() in EXTS:
            with Image.open(p) as im:
                im = im.convert("RGB")
                im.thumbnail(size)                    # en-boy oranını korur
                im.save(out / f"{p.stem}.jpg", "JPEG", quality=85)
                n += 1
    return n

if __name__ == "__main__":
    print(f"{thumbnails(Path(sys.argv[1]), Path('kucuk'))} images processed")
'''

PSUTIL_SYSINFO = r'''
# SİSTEM BİLGİSİ KALIBI — psutil, SABİT sayıda ölçüm (sonsuz döngü yok)
import time
from datetime import datetime
from pathlib import Path
import psutil

def snapshot() -> str:
    mem = psutil.virtual_memory()
    disk = psutil.disk_usage("/")
    return (f"{datetime.now():%Y-%m-%d %H:%M:%S} cpu={psutil.cpu_percent(interval=0.5)}% "
            f"ram={mem.percent}% ({mem.used // 2**20} MB) disk={disk.percent}%")

if __name__ == "__main__":
    lines = []
    for _ in range(3):
        lines.append(snapshot())
        time.sleep(1)
    Path("report.log").write_text("\n".join(lines) + "\n", encoding="utf-8")
'''

ZIP_BACKUP = r'''
# YEDEKLEME KALIBI — klasörü tarih damgalı zip'e, klasör yapısını koruyarak
import sys, zipfile
from datetime import datetime
from pathlib import Path

def backup(src: Path) -> Path:
    target = Path(f"yedek_{src.name}_{datetime.now():%Y%m%d_%H%M%S}.zip")   # çalışma klasörüne
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as zf:
        for p in sorted(src.rglob("*")):
            if p.is_file():
                zf.write(p, p.relative_to(src))
    return target

if __name__ == "__main__":
    z = backup(Path(sys.argv[1]))
    with zipfile.ZipFile(z) as zf:
        print(f"{z}: {len(zf.namelist())} dosya")
'''

REGEX_EXTRACT = r'''
# METİNDEN BİLGİ AYIKLAMA KALIBI — regex, tekrarları at, CSV'ye yaz
import csv, re, sys
from pathlib import Path

EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
PHONE = re.compile(r"(?:\+90\s?)?0?\(?5\d{2}\)?[\s-]?\d{3}[\s-]?\d{2}[\s-]?\d{2}")

def extract(text: str) -> list[tuple[str, str]]:
    found = [("email", m.group(0)) for m in EMAIL.finditer(text)]
    found += [("telefon", re.sub(r"\D", "", m.group(0))[-10:]) for m in PHONE.finditer(text)]
    return list(dict.fromkeys(found))               # sırayı koruyarak tekrarları at

if __name__ == "__main__":
    rows = extract(Path(sys.argv[1]).read_text(encoding="utf-8"))
    with open("bulunanlar.csv", "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(["tur", "deger"])
        writer.writerows(rows)
'''

HTTP_PARALLEL = r'''
# ÇOK SAYIDA ADRESİ KONTROL KALIBI — paralel, zaman aşımlı, sınırlı yeniden deneme
import sys
from concurrent.futures import ThreadPoolExecutor
import requests

HEADERS = {"User-Agent": "Mozilla/5.0"}

def status(url: str) -> tuple[str, str]:
    err = "?"
    for _ in range(2):
        try:
            r = requests.head(url, headers=HEADERS, timeout=10, allow_redirects=True)
            if r.status_code in (405, 501):            # HEAD desteklenmiyorsa GET
                r = requests.get(url, headers=HEADERS, timeout=10)
            return url, str(r.status_code)
        except requests.RequestException as e:
            err = type(e).__name__
    return url, f"HATA {err}"

if __name__ == "__main__":
    urls = [u.strip() for u in open(sys.argv[1], encoding="utf-8") if u.strip()]
    with ThreadPoolExecutor(max_workers=8) as ex:  # GUI yok; sonuçları beklemek doğru
        results = list(ex.map(status, urls))
    with open("linkler.txt", "w", encoding="utf-8") as fh:
        fh.writelines(f"{u} {s}\n" for u, s in results)
'''

FILE_ORGANIZE = r'''
# DOSYA DÜZENLEME KALIBI — uzantıya/tarihe göre alt klasörlere KOPYALA (kaynağa dokunma)
import shutil, sys
from datetime import datetime
from pathlib import Path

def organize(src: Path, dest: Path, by: str = "ext") -> int:
    n = 0
    for p in sorted(src.rglob("*")):
        if not p.is_file():
            continue
        if by == "ext":
            key = p.suffix.lower().lstrip(".") or "uzantisiz"
        else:
            key = datetime.fromtimestamp(p.stat().st_mtime).strftime("%Y-%m")
        (dest / key).mkdir(parents=True, exist_ok=True)
        shutil.copy2(p, dest / key / p.name)            # copy2 tarihleri de korur
        n += 1
    return n

if __name__ == "__main__":
    by = sys.argv[2] if len(sys.argv) > 2 else "ext"
    print(f"{organize(Path(sys.argv[1]), Path('duzenli'), by)} files copied")
'''

PPTX_REPORT = r'''
# SUNUM KALIBI — python-pptx ile başlık + madde slaytları
from pptx import Presentation

def build(rows: list[tuple[str, str]], path: str = "sunum.pptx") -> int:
    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = "Rapor"
    title.placeholders[1].text = f"{len(rows)} başlık"
    for head, body in rows:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = head
        slide.placeholders[1].text_frame.text = body
    prs.save(path)                                      # göreli yol → çalışma klasörü
    return len(prs.slides)
'''
